"""Exportação do catálogo em .xlsx, nas mesmas colunas que a importação lê
(categoria, item, valor, unidade, descrição, observação). Assim o arquivo
exportado pode ser editado e importado de volta."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONTE = "Inter"
COLUNAS = (
    ("CATEGORIA", 24),
    ("ITEM", 40),
    ("VALOR", 14),
    ("UNIDADE", 10),
    ("DESCRIÇÃO", 50),
    ("OBSERVAÇÃO", 36),
)
MOEDA = '"R$" #,##0.00'


def exportar_catalogo_xlsx(itens) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Catálogo"
    ws.sheet_view.showGridLines = False
    linha_fina = Border(bottom=Side(style="thin", color="E7E9EE"))
    for i, (titulo, largura) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=i, value=titulo)
        cel.font = Font(name=FONTE, size=8, bold=True, color="586070")
        cel.fill = PatternFill("solid", fgColor="F7F8FA")
        cel.border = Border(bottom=Side(style="thin", color="D8DBE2"))
        cel.alignment = Alignment(
            vertical="center", horizontal="right" if titulo == "VALOR" else "left"
        )
        ws.column_dimensions[cel.column_letter].width = largura
    ws.row_dimensions[1].height = 22

    iteraveis = itens.select_related("categoria") if hasattr(itens, "select_related") else itens
    linha = 2
    for item in iteraveis:
        valores = (
            item.categoria.nome if item.categoria else "",
            item.nome,
            float(item.valor_unitario_padrao),
            item.unidade_medida,
            item.descricao_padrao or "",
            item.observacoes or "",
        )
        for col, valor in enumerate(valores, start=1):
            cel = ws.cell(row=linha, column=col, value=valor)
            cel.font = Font(name=FONTE, size=10, color="111827", bold=col == 2)
            cel.border = linha_fina
            cel.alignment = Alignment(vertical="top", wrap_text=col in (2, 5, 6))
        ws.cell(row=linha, column=3).number_format = MOEDA
        linha += 1

    ws.freeze_panes = "A2"
    if linha > 2:
        ws.auto_filter.ref = f"A1:F{linha - 1}"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
