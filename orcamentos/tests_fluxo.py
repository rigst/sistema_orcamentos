"""Criação sem travas, edição de quantidade e valor na linha e cliente opcional."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase, override_settings
from django.urls import reverse
from docx import Document
from openpyxl import load_workbook

from catalogo.models import CategoriaItem, ItemCatalogo
from clientes.models import Cliente
from core.models import Empresa
from core.testing import SENHA_TESTE
from relatorios.exporters import (
    gerar_docx_memorial_descritivo,
    gerar_excel_orcamento,
    gerar_pdf_memorial_descritivo,
    gerar_pdf_orcamento,
    obter_alerta_status,
)
from relatorios.models import ConfiguracaoEmpresa

from .models import ItemOrcamento, Orcamento


class BaseFluxoTests(TestCase):
    def setUp(self):
        self.grupo = Group.objects.create(name="Empresa A")
        Empresa.objects.get_or_create(grupo=self.grupo, defaults={"nome": "Empresa A"})
        self.user = get_user_model().objects.create_user(
            username="orcamentista", password=SENHA_TESTE, perfil="orcamentista"
        )
        self.user.groups.add(self.grupo)
        self.client.force_login(self.user)
        self.cliente = Cliente.objects.create(nome_razao_social="Prefeitura", empresa=self.grupo)

    def novo_orcamento(self, **campos):
        dados = {
            "empresa": self.grupo,
            "titulo": "Edital",
            "data_emissao": date(2026, 10, 22),
            "criado_por": self.user,
            "atualizado_por": self.user,
        }
        dados.update(campos)
        return Orcamento.objects.create(**dados)


class CriacaoSemTravasTests(BaseFluxoTests):
    def test_tela_novo_mostra_os_dois_caminhos(self):
        with override_settings(ANTHROPIC_API_KEY="chave"):
            resposta = self.client.get(reverse("orcamentos:novo"))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, reverse("memoriais:novo_com_memorial"))
        self.assertContains(resposta, reverse("orcamentos:criar_rapido"))
        self.assertContains(resposta, reverse("orcamentos:criar"))

    @override_settings(ANTHROPIC_API_KEY="")
    def test_sem_ia_a_tela_novo_mostra_so_o_caminho_em_branco(self):
        resposta = self.client.get(reverse("orcamentos:novo"))
        self.assertNotContains(resposta, reverse("memoriais:novo_com_memorial"))
        self.assertContains(resposta, reverse("orcamentos:criar_rapido"))

    def test_criar_rapido_sem_nada_usa_titulo_generico_e_sem_cliente(self):
        resposta = self.client.post(reverse("orcamentos:criar_rapido"), {})
        orcamento = Orcamento.objects.get()
        self.assertRedirects(resposta, reverse("orcamentos:editar", args=[orcamento.pk]))
        self.assertEqual(orcamento.titulo, Orcamento.TITULO_GENERICO)
        self.assertIsNone(orcamento.cliente_id)
        self.assertEqual(orcamento.empresa, self.grupo)
        self.assertTrue(orcamento.numero.startswith("ORC-"))

    def test_criar_rapido_com_titulo_e_cliente(self):
        self.client.post(
            reverse("orcamentos:criar_rapido"),
            {"titulo": "  Estande Expointer  ", "cliente": self.cliente.pk},
        )
        orcamento = Orcamento.objects.get()
        self.assertEqual(orcamento.titulo, "Estande Expointer")
        self.assertEqual(orcamento.cliente, self.cliente)

    def test_criar_rapido_nao_aceita_cliente_de_outra_empresa(self):
        outro = Group.objects.create(name="Empresa B")
        alheio = Cliente.objects.create(nome_razao_social="Da B", empresa=outro)
        resposta = self.client.post(reverse("orcamentos:criar_rapido"), {"cliente": alheio.pk})
        self.assertEqual(resposta.status_code, 400)
        self.assertFalse(Orcamento.objects.exists())

    def test_formulario_completo_aceita_sem_cliente_e_sem_titulo(self):
        resposta = self.client.post(
            reverse("orcamentos:criar"),
            {
                "numero": "",
                "cliente": "",
                "titulo": "",
                "status": "rascunho",
                "data_emissao": "2026-10-22",
                "desconto_global_valor": "0,00",
                "desconto_global_percentual": "0,00",
                "acrescimo_global_valor": "0,00",
                "acrescimo_global_percentual": "0,00",
            },
        )
        orcamento = Orcamento.objects.get()
        self.assertRedirects(resposta, reverse("orcamentos:editar", args=[orcamento.pk]))
        self.assertEqual(orcamento.titulo, Orcamento.TITULO_GENERICO)

    def test_orcamento_sem_cliente_aparece_como_a_definir(self):
        orcamento = self.novo_orcamento()
        self.assertEqual(orcamento.cliente_exibicao, Orcamento.SEM_CLIENTE)
        self.assertIn(Orcamento.SEM_CLIENTE, str(orcamento))
        for url in (
            reverse("orcamentos:lista"),
            reverse("orcamentos:editar", args=[orcamento.pk]),
            reverse("relatorios:orcamento_central", args=[orcamento.pk]),
        ):
            resposta = self.client.get(url)
            self.assertEqual(resposta.status_code, 200, url)
            self.assertContains(resposta, Orcamento.SEM_CLIENTE)

    def test_todos_os_relatorios_saem_sem_cliente(self):
        orcamento = self.novo_orcamento()
        ItemOrcamento.objects.create(orcamento=orcamento, nome="Piso", valor_unitario=10, ordem=1)
        configuracao = ConfiguracaoEmpresa.objects.create(nome_empresa="Empresa A")
        status = obter_alerta_status(orcamento)

        pdf = gerar_pdf_orcamento(orcamento, configuracao, status)
        self.assertTrue(bytes(pdf).startswith(b"%PDF"))
        excel = load_workbook(BytesIO(gerar_excel_orcamento(orcamento, configuracao, status)))
        self.assertEqual(excel["Orçamento"]["A5"].value, Orcamento.SEM_CLIENTE)
        memorial_pdf = gerar_pdf_memorial_descritivo(orcamento, configuracao)
        self.assertTrue(bytes(memorial_pdf).startswith(b"%PDF"))
        memorial_doc = Document(BytesIO(gerar_docx_memorial_descritivo(orcamento, configuracao)))
        textos = [c.text for t in memorial_doc.tables for r in t.rows for c in r.cells]
        self.assertIn(Orcamento.SEM_CLIENTE, textos)


class EdicaoNaLinhaTests(BaseFluxoTests):
    def setUp(self):
        super().setUp()
        self.orcamento = self.novo_orcamento(cliente=self.cliente)
        self.item = ItemOrcamento.objects.create(
            orcamento=self.orcamento,
            nome="Piso",
            quantidade=Decimal("10"),
            valor_unitario=Decimal("50"),
            ordem=1,
        )
        self.outro = ItemOrcamento.objects.create(
            orcamento=self.orcamento,
            nome="Carpete",
            quantidade=Decimal("2"),
            valor_unitario=Decimal("100"),
            ordem=2,
        )
        self.url = reverse("orcamentos:item_edicao_rapida", args=[self.orcamento.pk, self.item.pk])

    def post(self, **dados):
        return self.client.post(self.url, dados, HTTP_X_REQUESTED_WITH="XMLHttpRequest")

    def test_muda_quantidade_e_valor_e_atualiza_o_total(self):
        resposta = self.post(quantidade="27,5", valor_unitario="1.234,56")
        self.assertEqual(resposta.status_code, 200)
        dados = resposta.json()

        self.item.refresh_from_db()
        self.orcamento.refresh_from_db()
        self.assertEqual(self.item.quantidade, Decimal("27.50"))
        self.assertEqual(self.item.valor_unitario, Decimal("1234.56"))
        esperado_item = Decimal("27.50") * Decimal("1234.56")
        self.assertEqual(self.item.subtotal, esperado_item.quantize(Decimal("0.01")))
        self.assertEqual(self.orcamento.subtotal_itens, self.item.subtotal + Decimal("200.00"))
        self.assertEqual(dados["subtotal"], "R$ 33.950,40")
        self.assertEqual(dados["total_final"], "R$ 34.150,40")
        self.assertIn("34.150,40", dados["totais_html"])

    def test_aceita_so_um_dos_campos(self):
        self.post(valor_unitario="60")
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantidade, Decimal("10.00"))
        self.assertEqual(self.item.valor_unitario, Decimal("60.00"))

    def test_respeita_desconto_e_acrescimo_do_item(self):
        self.item.desconto_percentual = Decimal("10")
        self.item.save()
        dados = self.post(quantidade="10", valor_unitario="100").json()
        self.assertEqual(dados["subtotal"], "R$ 900,00")

    def test_numero_invalido_nao_salva(self):
        for valor in ("abc", "", "1e400", "99999999999"):
            resposta = self.post(quantidade=valor, valor_unitario="10")
            self.assertEqual(resposta.status_code, 400, valor)
            self.assertIn("erro", resposta.json())
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantidade, Decimal("10.00"))

    def test_quantidade_zero_e_recusada_pelo_modelo(self):
        resposta = self.post(quantidade="0", valor_unitario="10")
        self.assertEqual(resposta.status_code, 400)

    def test_exige_post(self):
        self.assertEqual(self.client.get(self.url).status_code, 405)

    def test_orcamento_fora_de_rascunho_nao_edita(self):
        self.orcamento.status = "enviado"
        self.orcamento.save()
        self.assertEqual(self.post(quantidade="5").status_code, 403)

    def test_item_de_outro_orcamento_da_404(self):
        outro_orcamento = self.novo_orcamento(titulo="Outro")
        url = reverse("orcamentos:item_edicao_rapida", args=[outro_orcamento.pk, self.item.pk])
        resposta = self.client.post(url, {"quantidade": "5"})
        self.assertEqual(resposta.status_code, 404)

    def test_orcamento_de_outra_empresa_da_404(self):
        outro = Group.objects.create(name="Empresa B")
        alheio = self.novo_orcamento(empresa=outro)
        item = ItemOrcamento.objects.create(orcamento=alheio, nome="X", valor_unitario=1, ordem=1)
        url = reverse("orcamentos:item_edicao_rapida", args=[alheio.pk, item.pk])
        self.assertEqual(self.client.post(url, {"quantidade": "5"}).status_code, 404)

    def test_tabela_traz_campos_editaveis_e_o_script(self):
        resposta = self.client.get(reverse("orcamentos:editar", args=[self.orcamento.pk]))
        self.assertContains(resposta, f'data-linha-rapida="{self.url}"')
        self.assertContains(resposta, 'data-campo-rapido="quantidade"', count=2)
        self.assertContains(resposta, 'data-campo-rapido="valor_unitario"', count=2)
        self.assertContains(resposta, "js/orcamento-linhas.js")

    def test_somente_leitura_nao_traz_campos(self):
        self.orcamento.status = "enviado"
        self.orcamento.save()
        resposta = self.client.get(reverse("orcamentos:editar", args=[self.orcamento.pk]))
        self.assertNotContains(resposta, "data-campo-rapido")
        self.assertNotContains(resposta, "js/orcamento-linhas.js")

    def test_item_do_catalogo_editado_na_linha_fica_marcado_como_personalizado(self):
        categoria = CategoriaItem.objects.create(nome="Pisos", empresa=self.grupo)
        catalogo = ItemCatalogo.objects.create(
            nome="Piso",
            categoria=categoria,
            unidade_medida="un",
            valor_unitario_padrao=Decimal("50"),
            empresa=self.grupo,
        )
        self.item.item_catalogo = catalogo
        self.item.save()
        dados = self.post(valor_unitario="70").json()
        self.assertTrue(dados["diverge_catalogo"])
        catalogo.refresh_from_db()
        self.assertEqual(catalogo.valor_unitario_padrao, Decimal("50.00"))
