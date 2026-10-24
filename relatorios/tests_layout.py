"""Garantias de layout dos relatórios novos (PDF com as fontes do site, números
sem quebra, planilha com números de verdade e Word editável)."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from docx import Document
from openpyxl import load_workbook

from catalogo.models import CategoriaItem, ItemCatalogo
from clientes.models import Cliente
from orcamentos.models import ItemOrcamento, Orcamento

from . import exporters as exp
from .conteudo import formatar_moeda
from .models import ConfiguracaoEmpresa


class RelatoriosLayoutTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        empresa = Group.objects.create(name="Empresa layout")
        usuario = get_user_model().objects.create_user(username="layout")
        cliente = Cliente.objects.create(nome_razao_social="Cliente Layout", empresa=empresa)
        cls.orcamento = Orcamento.objects.create(
            titulo="Feira de teste",
            cliente=cliente,
            empresa=empresa,
            criado_por=usuario,
            atualizado_por=usuario,
            evento_nome="Feira",
        )
        categorias = [
            CategoriaItem.objects.create(nome=f"Categoria {n}", empresa=empresa, cor="#2563EB")
            for n in range(3)
        ]
        catalogo = [
            ItemCatalogo.objects.create(
                nome=f"Base {n}", categoria=categorias[n], empresa=empresa, valor_unitario_padrao=1
            )
            for n in range(3)
        ]
        for n in range(40):
            ItemOrcamento.objects.create(
                orcamento=cls.orcamento,
                item_catalogo=catalogo[n % 3],
                nome=f"Serviço número {n} com um nome razoavelmente comprido para quebrar",
                descricao="Descrição longa " * 8,
                quantidade=Decimal("12.5"),
                unidade_medida="m2",
                valor_unitario=Decimal("38000.00"),
            )
        cls.orcamento.refresh_from_db()
        cls.configuracao = ConfiguracaoEmpresa.objects.create(
            empresa=empresa, nome_empresa="Empresa Layout Ltda.", email="a@b.com"
        )

    def test_pdf_usa_as_fontes_do_site_embutidas(self):
        conteudo = exp.gerar_pdf_orcamento(
            self.orcamento, self.configuracao, exp.obter_alerta_status(self.orcamento)
        )
        self.assertIn(b"Inter-Regular", conteudo)
        self.assertIn(b"Manrope", conteudo)
        self.assertIn(b"FontFile2", conteudo)  # TrueType embutida, não referência

    def test_pdf_compacto_para_orcamento_grande(self):
        conteudo = exp.gerar_pdf_orcamento(
            self.orcamento, self.configuracao, exp.obter_alerta_status(self.orcamento)
        )
        paginas = conteudo.count(b"/Type /Page\n") or conteudo.count(b"/Type /Page ")
        # 40 itens com descrição longa cabiam em ~9 páginas no layout antigo.
        self.assertLessEqual(paginas, 6)

    def test_valores_monetarios_nunca_quebram_linha(self):
        from reportlab.platypus import Paragraph

        from .pdf_estilo import estilos, registrar_fontes

        registrar_fontes()
        est = estilos()
        maior = formatar_moeda(Decimal("12.5") * Decimal("38000"))
        p = Paragraph(maior.replace(" ", "&nbsp;"), est["celula_num_forte"])
        _, altura = p.wrap(29 * 2.8346, 1000)  # largura da coluna SUBTOTAL, em pt
        self.assertLess(altura, est["celula_num_forte"].leading * 1.5)

    def test_planilha_com_numeros_formatados_e_cabecalho_fixo(self):
        conteudo = exp.gerar_excel_orcamento(
            self.orcamento, self.configuracao, exp.obter_alerta_status(self.orcamento)
        )
        planilha = load_workbook(BytesIO(conteudo))
        aba = planilha["Orçamento"]
        self.assertEqual(aba.freeze_panes, "A8")
        primeira = [c.value for c in aba[8]]
        self.assertEqual(primeira[12], 475000.0)
        self.assertEqual(aba.cell(row=8, column=13).number_format, '"R$" #,##0.00')
        categorias = planilha["Por categoria"]
        self.assertEqual(categorias.cell(row=5, column=3).value, "=SUM(C2:C4)")

    def test_memorial_docx_editavel_com_fonte_do_site(self):
        conteudo = exp.gerar_docx_memorial_descritivo(self.orcamento, self.configuracao)
        documento = Document(BytesIO(conteudo))
        self.assertEqual(documento.styles["Normal"].font.name, "Inter")
        textos = "\n".join(p.text for p in documento.paragraphs)
        self.assertIn("Categoria 0", textos)
        self.assertIn("Empresa Layout Ltda.", textos)
        self.assertGreaterEqual(textos.count("12,50 metros quadrados"), 40)
