from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from docx import Document
from openpyxl import load_workbook
from PIL import Image

from catalogo.models import CategoriaItem, ItemCatalogo
from clientes.models import Cliente
from core.testing import SENHA_TESTE
from orcamentos.models import ItemOrcamento, Orcamento

from .formatos import DOCX, XLSX
from .models import ConfiguracaoEmpresa


def texto_da_planilha(planilha) -> str:
    return "\n".join(
        str(celula)
        for aba in planilha.worksheets
        for linha in aba.iter_rows(values_only=True)
        for celula in linha
        if celula is not None
    )


def texto_do_docx(conteudo: bytes) -> str:
    documento = Document(BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    for tabela in documento.tables:
        partes += [celula.text for linha in tabela.rows for celula in linha.cells]
    return "\n".join(partes)


class RelatoriosPermissaoTests(TestCase):
    def gerar_logo_teste(self):
        buffer = BytesIO()
        Image.new("RGB", (180, 60), color=(80, 140, 220)).save(buffer, format="PNG")
        return SimpleUploadedFile("logo.png", buffer.getvalue(), content_type="image/png")

    def test_orcamentista_nao_pode_alterar_dados_da_empresa(self):
        user = get_user_model().objects.create_user(
            username="orcamentista",
            password=SENHA_TESTE,
            perfil="orcamentista",
        )
        self.client.force_login(user)

        response = self.client.post(reverse("relatorios:minha_empresa"), {"nome_empresa": "X"})

        self.assertIn(response.status_code, (302, 403))
        self.assertFalse(ConfiguracaoEmpresa.objects.filter(nome_empresa="X").exists())

    def test_visualizador_ve_os_dados_da_propria_empresa(self):
        user = get_user_model().objects.create_user(
            username="visualizador_empresa",
            password=SENHA_TESTE,
            perfil="visualizador",
        )
        ConfiguracaoEmpresa.objects.create(nome_empresa="Empresa Visivel")
        self.client.force_login(user)

        response = self.client.get(reverse("relatorios:minha_empresa"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Empresa Visivel")

    def test_visualizador_nao_pode_alterar_opcoes_do_relatorio(self):
        user = get_user_model().objects.create_user(
            username="visualizador_relatorio",
            password=SENHA_TESTE,
            perfil="visualizador",
        )
        cliente = Cliente.objects.create(nome_razao_social="Cliente Permissao")
        orcamento = Orcamento.objects.create(
            numero="ORC-REL-0009",
            cliente=cliente,
            titulo="Teste permissao",
            status="rascunho",
            data_emissao="2026-04-09",
            criado_por=user,
            atualizado_por=user,
        )
        self.client.force_login(user)

        response = self.client.post(
            reverse("relatorios:orcamento_central", args=[orcamento.pk]),
            {"mostrar_ajustes_no_relatorio": "on"},
        )

        self.assertEqual(response.status_code, 403)

    def test_logo_configuracao_exige_login(self):
        configuracao = ConfiguracaoEmpresa.objects.create(
            nome_empresa="Empresa Visivel", logo=self.gerar_logo_teste()
        )

        response = self.client.get(reverse("relatorios:configuracao_logo", args=[configuracao.pk]))

        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response.url)

    def test_logo_configuracao_so_e_acessivel_na_mesma_empresa(self):
        empresa_a = Group.objects.create(name="Empresa A")
        empresa_b = Group.objects.create(name="Empresa B")
        usuario_a = get_user_model().objects.create_user(
            username="usuario_logo_a",
            password=SENHA_TESTE,
            perfil="visualizador",
        )
        usuario_a.groups.set([empresa_a])
        usuario_b = get_user_model().objects.create_user(
            username="usuario_logo_b",
            password=SENHA_TESTE,
            perfil="visualizador",
        )
        usuario_b.groups.set([empresa_b])

        configuracao_a = ConfiguracaoEmpresa.objects.create(
            nome_empresa="Empresa A",
            empresa=empresa_a,
            logo=self.gerar_logo_teste(),
        )
        configuracao_b = ConfiguracaoEmpresa.objects.create(
            nome_empresa="Empresa B",
            empresa=empresa_b,
            logo=self.gerar_logo_teste(),
        )

        self.client.force_login(usuario_a)
        response_propria = self.client.get(
            reverse("relatorios:configuracao_logo", args=[configuracao_a.pk])
        )
        response_outra = self.client.get(
            reverse("relatorios:configuracao_logo", args=[configuracao_b.pk])
        )
        self.client.force_login(usuario_b)
        response_b = self.client.get(
            reverse("relatorios:configuracao_logo", args=[configuracao_b.pk])
        )

        self.assertEqual(response_propria.status_code, 200)
        self.assertEqual(response_propria["Content-Type"], "image/png")
        self.assertEqual(response_outra.status_code, 404)
        self.assertEqual(response_b.status_code, 200)


class RelatoriosExportacaoTests(TestCase):
    def gerar_logo_teste(self):
        buffer = BytesIO()
        Image.new("RGB", (180, 60), color=(80, 140, 220)).save(buffer, format="PNG")
        return SimpleUploadedFile("logo.png", buffer.getvalue(), content_type="image/png")

    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="orcamentista_export",
            password=SENHA_TESTE,
            perfil="orcamentista",
        )
        self.client.force_login(self.user)

        self.cliente = Cliente.objects.create(nome_razao_social="Cliente Exportacao")
        self.orcamento = Orcamento.objects.create(
            numero="ORC-REL-0001",
            cliente=self.cliente,
            titulo="Relatorio profissional",
            status="rascunho",
            data_emissao="2026-04-09",
            criado_por=self.user,
            atualizado_por=self.user,
        )
        ItemOrcamento.objects.create(
            orcamento=self.orcamento,
            ordem=1,
            codigo_item="REL-1",
            nome="Servico relatorio",
            unidade_medida="un",
            quantidade=Decimal("1"),
            valor_unitario=Decimal("15000.00"),
        )
        self.orcamento.refresh_from_db()
        ConfiguracaoEmpresa.objects.create(
            nome_empresa="Empresa Exemplo",
            logo=self.gerar_logo_teste(),
        )

    def test_central_de_relatorio_mostra_alerta_de_status(self):
        response = self.client.get(
            reverse("relatorios:orcamento_central", args=[self.orcamento.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Orçamento em rascunho")
        self.assertContains(response, "Logo da empresa")
        self.assertContains(
            response,
            reverse("relatorios:configuracao_logo", args=[ConfiguracaoEmpresa.objects.get().pk]),
        )
        self.assertContains(response, "R$ 15.000,00")

    def test_central_de_relatorio_atualiza_opcao_de_mostrar_ajustes(self):
        response = self.client.post(
            reverse("relatorios:orcamento_central", args=[self.orcamento.pk]),
            {"mostrar_ajustes_no_relatorio": "on"},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.orcamento.refresh_from_db()
        self.assertTrue(self.orcamento.mostrar_ajustes_no_relatorio)
        self.assertContains(response, "Opções do relatório atualizadas.")

    def test_central_de_relatorio_exibe_opcoes_adicionais_e_exportacao_word(self):
        response = self.client.get(
            reverse("relatorios:orcamento_central", args=[self.orcamento.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Mostrar descrição inicial")
        self.assertContains(response, "Mostrar informações financeiras")
        self.assertContains(
            response, reverse("relatorios:orcamento_memorial_word", args=[self.orcamento.pk])
        )
        self.assertContains(response, "Baixar Memorial em Word")

    def test_relatorios_respeitam_configuracao_vinculada_ao_orcamento(self):
        self.orcamento.configuracao_empresa = ConfiguracaoEmpresa.objects.create(
            nome_empresa="Empresa Escolhida",
            cidade="Curitiba",
        )
        self.orcamento.save(update_fields=["configuracao_empresa", "atualizado_em"])
        ConfiguracaoEmpresa.objects.create(
            nome_empresa="Empresa Mais Recente",
            cidade="Campinas",
        )

        response = self.client.get(
            reverse("relatorios:orcamento_central", args=[self.orcamento.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Empresa Escolhida")
        self.assertContains(response, "Curitiba")
        self.assertNotContains(response, "Empresa Mais Recente")

    def test_central_de_relatorio_atualiza_multiplas_opcoes(self):
        response = self.client.post(
            reverse("relatorios:orcamento_central", args=[self.orcamento.pk]),
            {
                "mostrar_ajustes_no_relatorio": "on",
                "mostrar_descricao_inicial_no_relatorio": "on",
                "mostrar_rodape_institucional_no_relatorio": "on",
                "mostrar_financeiro_no_memorial": "on",
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.orcamento.refresh_from_db()
        self.assertTrue(self.orcamento.mostrar_ajustes_no_relatorio)
        self.assertTrue(self.orcamento.mostrar_descricao_inicial_no_relatorio)
        self.assertFalse(self.orcamento.mostrar_observacoes_gerais_no_relatorio)
        self.assertTrue(self.orcamento.mostrar_rodape_institucional_no_relatorio)
        self.assertFalse(self.orcamento.mostrar_contatos_evento_no_memorial)
        self.assertTrue(self.orcamento.mostrar_financeiro_no_memorial)
        self.assertFalse(self.orcamento.mostrar_dados_contratuais_no_memorial)
        self.assertFalse(self.orcamento.mostrar_informacoes_complementares_no_memorial)
        self.assertContains(response, "Opções do relatório atualizadas.")

    def test_exportacao_excel_retorna_arquivo(self):
        categoria = CategoriaItem.objects.create(nome="Civil", cor="#2563EB")
        item_catalogo = ItemCatalogo.objects.create(
            codigo="CIV-EXCEL-01",
            nome="Servico civil",
            categoria=categoria,
            unidade_medida="sv",
            valor_unitario_padrao=Decimal("100.00"),
            descricao_padrao="Descricao do item no excel",
        )
        item = self.orcamento.itens.first()
        item.item_catalogo = item_catalogo
        item.descricao = "Descricao detalhada"
        item.save(update_fields=["item_catalogo", "descricao", "atualizado_em"])
        self.orcamento.descricao_inicial = "Introducao do PDF"
        self.orcamento.observacoes_gerais = "Observacoes do documento"
        self.orcamento.save(
            update_fields=["descricao_inicial", "observacoes_gerais", "atualizado_em"]
        )
        configuracao = ConfiguracaoEmpresa.objects.get()
        configuracao.rodape_relatorio = "Rodape institucional"
        configuracao.save(update_fields=["rodape_relatorio", "atualizado_em"])

        response = self.client.get(reverse("relatorios:orcamento_excel", args=[self.orcamento.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX)
        self.assertIn('.xlsx"', response["Content-Disposition"])
        planilha = load_workbook(BytesIO(response.content))
        textos = texto_da_planilha(planilha)
        self.assertIn("Orçamento em rascunho", textos)
        self.assertIn("CATEGORIA", textos)
        self.assertIn("Civil", textos)
        self.assertIn("Descricao detalhada", textos)
        self.assertIn("Introducao do PDF", textos)
        self.assertIn("Observacoes do documento", textos)
        self.assertIn("Rodape institucional", textos)
        # números são números (dá para somar no Excel), com formato de moeda
        aba = planilha["Orçamento"]
        subtotais = [
            linha
            for linha in aba.iter_rows(min_row=8, max_col=13)
            if isinstance(linha[12].value, (int, float)) and linha[3].value
        ]
        self.assertTrue(subtotais)
        self.assertEqual(subtotais[0][12].number_format, '"R$" #,##0.00')
        self.assertIn("Por categoria", planilha.sheetnames)

    def test_exportacao_pdf_retorna_arquivo(self):
        response = self.client.get(reverse("relatorios:orcamento_pdf", args=[self.orcamento.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-1.4"))
        self.assertIn(b"/Subtype /Image", response.content)

    def test_exportacao_memorial_descritivo_retorna_arquivo(self):
        categoria = CategoriaItem.objects.create(nome="Civil", cor="#2563EB")
        item_catalogo = ItemCatalogo.objects.create(
            codigo="CIV-01",
            nome="Servico civil",
            categoria=categoria,
            unidade_medida="sv",
            valor_unitario_padrao=Decimal("100.00"),
        )
        item = self.orcamento.itens.first()
        item.item_catalogo = item_catalogo
        item.descricao = "Descricao longa do servico"
        item.save()
        self.orcamento.evento_nome = "Expodireto 2026"
        self.orcamento.evento_local = "Nao-Me-Toque/RS"
        self.orcamento.evento_contato = "Mariana"
        self.orcamento.condicoes_pagamento = "30% na aprovacao"
        self.orcamento.contrato_razao_social = "Cliente Exportacao LTDA"
        self.orcamento.save(
            update_fields=[
                "evento_nome",
                "evento_local",
                "evento_contato",
                "condicoes_pagamento",
                "contrato_razao_social",
            ]
        )
        configuracao = ConfiguracaoEmpresa.objects.get()
        configuracao.dados_bancarios = "Banco 001 Ag 1234 Cc 99999-9"
        configuracao.chave_pix = "pix@empresa.com"
        configuracao.assinatura_nome = "Diretoria Comercial"
        configuracao.texto_institucional_memorial = "Documento emitido para fins comerciais."
        configuracao.save()

        response = self.client.get(
            reverse("relatorios:orcamento_memorial_pdf", args=[self.orcamento.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-1.4"))
        self.assertIn(b"Memorial Descritivo", response.content)  # título do documento
        # O texto do PDF sai das fontes embutidas (codificado); o conteúdo é o
        # mesmo do Word, conferido no .docx.
        word = self.client.get(
            reverse("relatorios:orcamento_memorial_word", args=[self.orcamento.pk])
        )
        textos = texto_do_docx(word.content)
        for esperado in (
            "Civil",
            "Descricao longa do servico",
            "Expodireto 2026",
            "30% na aprovacao",
            "Cliente Exportacao LTDA",
            "pix@empresa.com",
            "Documento emitido para fins comerciais.",
        ):
            self.assertIn(esperado, textos)

    def test_exportacao_memorial_word_retorna_arquivo(self):
        self.orcamento.evento_contato = "Mariana"
        self.orcamento.condicoes_pagamento = "30% na aprovacao"
        self.orcamento.contrato_razao_social = "Cliente Exportacao LTDA"
        self.orcamento.observacoes_gerais = "Observacoes para Word"
        self.orcamento.save(
            update_fields=[
                "evento_contato",
                "condicoes_pagamento",
                "contrato_razao_social",
                "observacoes_gerais",
                "atualizado_em",
            ]
        )
        configuracao = ConfiguracaoEmpresa.objects.get()
        configuracao.chave_pix = "pix@empresa.com"
        configuracao.save(update_fields=["chave_pix", "atualizado_em"])

        response = self.client.get(
            reverse("relatorios:orcamento_memorial_word", args=[self.orcamento.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], DOCX)
        self.assertIn('.docx"', response["Content-Disposition"])
        textos = texto_do_docx(response.content)
        self.assertIn("Relatorio profissional", textos)
        self.assertIn("Mariana", textos)
        self.assertIn("Cliente Exportacao LTDA", textos)
        self.assertIn("pix@empresa.com", textos)

    def test_exportacao_memorial_word_omite_blocos_desmarcados(self):
        self.orcamento.evento_contato = "Mariana"
        self.orcamento.condicoes_pagamento = "30% na aprovacao"
        self.orcamento.contrato_razao_social = "Cliente Exportacao LTDA"
        self.orcamento.observacoes_gerais = "Observacoes ocultas"
        self.orcamento.mostrar_contatos_evento_no_memorial = False
        self.orcamento.mostrar_financeiro_no_memorial = False
        self.orcamento.mostrar_dados_contratuais_no_memorial = False
        self.orcamento.mostrar_observacoes_gerais_no_relatorio = False
        self.orcamento.save(
            update_fields=[
                "evento_contato",
                "condicoes_pagamento",
                "contrato_razao_social",
                "observacoes_gerais",
                "mostrar_contatos_evento_no_memorial",
                "mostrar_financeiro_no_memorial",
                "mostrar_dados_contratuais_no_memorial",
                "mostrar_observacoes_gerais_no_relatorio",
                "atualizado_em",
            ]
        )

        response = self.client.get(
            reverse("relatorios:orcamento_memorial_word", args=[self.orcamento.pk])
        )

        self.assertEqual(response.status_code, 200)
        textos = texto_do_docx(response.content)
        self.assertNotIn("Mariana", textos)
        self.assertNotIn("Cliente Exportacao LTDA", textos)
        self.assertNotIn("30% na aprovacao", textos)
        self.assertNotIn("Observacoes ocultas", textos)


class RelatoriosValidacaoTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="admin_relatorios",
            password=SENHA_TESTE,
            perfil="admin",
        )
        self.client.force_login(self.user)

    def test_configuracao_empresa_invalida_exibe_erros(self):
        response = self.client.post(
            reverse("relatorios:minha_empresa"),
            {
                "nome_empresa": "",
                "cpf_cnpj": "123",
                "telefone": "9999",
                "cep": "123",
                "estado": "XX",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Informe o nome da empresa.")
        self.assertContains(response, "Informe um CPF com 11 dígitos ou um CNPJ com 14 dígitos.")
        self.assertContains(response, "Telefone deve ter 10 ou 11 dígitos, incluindo DDD.")
        self.assertContains(response, "CEP deve ter 8 dígitos.")
        self.assertContains(response, "Informe uma UF válida com duas letras.")

    def test_configuracao_empresa_aceita_campos_com_mascara(self):
        response = self.client.post(
            reverse("relatorios:minha_empresa"),
            {
                "nome_empresa": "Empresa Formatada",
                "cpf_cnpj": "12.345.678/0001-90",
                "telefone": "(11) 4002-8922",
                "cep": "01310-100",
                "estado": "SP",
                "dados_bancarios": "Banco Exemplo",
                "chave_pix": "financeiro@empresa.com",
                "validade_padrao_proposta": "15 dias",
                "assinatura_nome": "Diretoria",
                "assinatura_cargo": "Comercial",
                "assinatura_contato": "(11) 4002-8922",
                "texto_institucional_memorial": "Texto padrao",
            },
        )

        self.assertEqual(response.status_code, 302)
        configuracao = ConfiguracaoEmpresa.objects.get(nome_empresa="Empresa Formatada")
        self.assertEqual(configuracao.cpf_cnpj, "12.345.678/0001-90")
        self.assertEqual(configuracao.telefone, "(11) 4002-8922")
        self.assertEqual(configuracao.cep, "01310-100")
        self.assertEqual(configuracao.chave_pix, "financeiro@empresa.com")
        self.assertEqual(configuracao.validade_padrao_proposta, "15 dias")


class MinhaEmpresaTests(TestCase):
    """Um cadastro só por empresa, sem lista nem seletor."""

    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="admin_minha_empresa",
            password=SENHA_TESTE,
            perfil="admin",
        )
        self.client.force_login(self.user)

    def _dados(self, **extra):
        return {"nome_empresa": "Empresa Unica", **extra}

    def test_primeiro_salvamento_cria_e_os_seguintes_editam_o_mesmo_cadastro(self):
        url = reverse("relatorios:minha_empresa")
        self.client.post(url, self._dados())
        self.client.post(url, self._dados(nome_empresa="Empresa Renomeada"))

        self.assertEqual(ConfiguracaoEmpresa.objects.count(), 1)
        self.assertEqual(ConfiguracaoEmpresa.objects.get().nome_empresa, "Empresa Renomeada")

    def test_cadastro_novo_passa_a_valer_para_orcamentos_que_estavam_sem(self):
        cliente = Cliente.objects.create(nome_razao_social="Cliente X")
        orcamento = Orcamento.objects.create(
            cliente=cliente, titulo="Sem cadastro", criado_por=self.user, atualizado_por=self.user
        )
        Orcamento.objects.filter(pk=orcamento.pk).update(configuracao_empresa=None)

        self.client.post(reverse("relatorios:minha_empresa"), self._dados())

        orcamento.refresh_from_db()
        self.assertEqual(orcamento.configuracao_empresa.nome_empresa, "Empresa Unica")

    def test_padroes_comerciais_entram_no_orcamento_novo(self):
        self.client.post(
            reverse("relatorios:minha_empresa"),
            self._dados(
                condicoes_pagamento_padrao="50% na assinatura",
                servicos_taxas_inclusos_padrao="Montagem e frete",
            ),
        )
        orcamento = Orcamento.objects.create(
            titulo="Novo", criado_por=self.user, atualizado_por=self.user
        )

        self.assertEqual(orcamento.condicoes_pagamento, "50% na assinatura")
        self.assertEqual(orcamento.servicos_taxas_inclusos, "Montagem e frete")

    def test_rotas_antigas_de_configuracao_levam_para_minha_empresa(self):
        destino = reverse("relatorios:minha_empresa")
        for rota in ("/relatorios/configuracoes/", "/relatorios/configuracoes/nova/"):
            self.assertRedirects(self.client.get(rota), destino)

    def test_admin_nao_permite_criar_segundo_cadastro(self):
        self.user.is_staff = True
        self.user.save()
        response = self.client.get(reverse("admin:relatorios_configuracaoempresa_add"))
        self.assertEqual(response.status_code, 403)


class RelatoriosAtualizacaoTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="admin_edita_empresa",
            password=SENHA_TESTE,
            perfil="admin",
        )
        self.client.force_login(self.user)

    def test_formulario_de_edicao_exibe_aviso_sobre_orcamentos_nao_enviados(self):
        configuracao = ConfiguracaoEmpresa.objects.create(nome_empresa="Empresa Base")

        response = self.client.get(reverse("relatorios:minha_empresa"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, configuracao.nome_empresa)
        self.assertContains(response, "orçamentos ainda não enviados")
        self.assertContains(response, "Salvar empresa")

    def test_edicao_de_configuracao_exibe_mensagem_de_propagacao(self):
        configuracao = ConfiguracaoEmpresa.objects.create(nome_empresa="Empresa Base")

        response = self.client.post(
            reverse("relatorios:minha_empresa"),
            {
                "nome_empresa": "Empresa Atualizada",
                "nome_fantasia": configuracao.nome_fantasia,
                "cpf_cnpj": configuracao.cpf_cnpj,
                "email": configuracao.email,
                "telefone": configuracao.telefone,
                "site": configuracao.site,
                "cep": configuracao.cep,
                "logradouro": configuracao.logradouro,
                "numero": configuracao.numero,
                "complemento": configuracao.complemento,
                "bairro": configuracao.bairro,
                "cidade": configuracao.cidade,
                "estado": configuracao.estado,
                "rodape_relatorio": configuracao.rodape_relatorio,
                "ativo": "on",
            },
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "Os relatórios e os orçamentos ainda não enviados passam a usar essas informações.",
        )
