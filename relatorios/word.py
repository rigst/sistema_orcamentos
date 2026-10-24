"""Memorial descritivo em .docx (python-docx), editável no Word e no Google
Docs, com a mesma estrutura e paleta do PDF."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import conteudo as c

FONTE = "Inter"
TINTA = RGBColor(0x11, 0x18, 0x27)
TEXTO = RGBColor(0x37, 0x41, 0x51)
SUAVE = RGBColor(0x58, 0x60, 0x70)


def _run(paragrafo, texto: str, *, tamanho=10.5, negrito=False, cor=TEXTO, fonte=FONTE):
    run = paragrafo.add_run(texto)
    run.font.name = fonte
    run._element.rPr.rFonts.set(qn("w:eastAsia"), fonte)
    run.font.size = Pt(tamanho)
    run.font.bold = negrito
    run.font.color.rgb = cor
    return run


def _paragrafo(doc, texto: str = "", *, depois=6, antes=0, alinhamento=None, **kw):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(depois)
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.line_spacing = 1.25
    if alinhamento is not None:
        p.alignment = alinhamento
    if texto:
        _run(p, texto, **kw)
    return p


def _rotulo_secao(doc, titulo: str):
    p = _paragrafo(doc, depois=6, antes=14)
    _run(p, titulo.upper(), tamanho=8, negrito=True, cor=SUAVE)
    return p


def _linha_inferior(paragrafo, cor="D8DBE2", espessura=6):
    pPr = paragrafo._p.get_or_add_pPr()
    borda = OxmlElement("w:pBdr")
    baixo = OxmlElement("w:bottom")
    for chave, valor in (
        ("w:val", "single"),
        ("w:sz", str(espessura)),
        ("w:space", "4"),
        ("w:color", cor),
    ):
        baixo.set(qn(chave), valor)
    borda.append(baixo)
    pPr.append(borda)


def _sem_bordas(tabela):
    tbl = tabela._tbl
    tblPr = tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "nil")
        bordas.append(el)
    h = OxmlElement("w:insideH")
    for chave, valor in (("w:val", "single"), ("w:sz", "4"), ("w:color", "E7E9EE")):
        h.set(qn(chave), valor)
    bordas.append(h)
    tblPr.append(bordas)


def _tabela_campos(doc, campos: list[c.Campo], largura_rotulo=Cm(4.6)):
    tabela = doc.add_table(rows=0, cols=2)
    tabela.alignment = WD_TABLE_ALIGNMENT.LEFT
    _sem_bordas(tabela)
    for campo in campos:
        rot, val = tabela.add_row().cells
        rot.width, val.width = largura_rotulo, Cm(17.6) - largura_rotulo
        rot.paragraphs[0].paragraph_format.space_after = Pt(3)
        val.paragraphs[0].paragraph_format.space_after = Pt(3)
        _run(rot.paragraphs[0], campo.rotulo, tamanho=9, cor=SUAVE)
        _run(val.paragraphs[0], campo.valor, tamanho=10.5, cor=TINTA)
    return tabela


def _pagina(doc) -> None:
    secao = doc.sections[0]
    secao.page_height, secao.page_width = Cm(29.7), Cm(21.0)
    secao.left_margin = secao.right_margin = Cm(1.7)
    secao.top_margin, secao.bottom_margin = Cm(1.6), Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(10.5)


def _topo(doc, orcamento, configuracao, memorial) -> None:
    empresa = c.nome_empresa(configuracao)
    if empresa:
        p = _paragrafo(doc, depois=2)
        _run(p, empresa, tamanho=13, negrito=True, cor=TINTA, fonte="Manrope")
    contato = c.contato_empresa(configuracao)
    if empresa or contato:
        _linha_inferior(_paragrafo(doc, " · ".join(contato), depois=14, tamanho=8.5, cor=SUAVE))
    _paragrafo(
        doc,
        f"MEMORIAL DESCRITIVO · {orcamento.numero}",
        depois=2,
        antes=6,
        tamanho=8,
        negrito=True,
        cor=SUAVE,
    )
    _paragrafo(
        doc,
        orcamento.evento_nome or orcamento.titulo,
        depois=2,
        tamanho=22,
        negrito=True,
        cor=TINTA,
        fonte="Manrope",
    )
    _paragrafo(doc, memorial.local_data, depois=12, tamanho=11, cor=SUAVE)


def _categorias(doc, categorias) -> None:
    for grupo in categorias:
        p = _paragrafo(doc, depois=6, antes=14)
        cor = (grupo["categoria_cor"] or "#94A3B8").lstrip("#")
        _run(p, "● ", tamanho=10, cor=RGBColor.from_string(cor.upper()))
        _run(p, grupo["categoria_nome"], tamanho=11, negrito=True, cor=TINTA)
        _run(p, f"\t{c.formatar_moeda(grupo['subtotal'])}", tamanho=11, negrito=True, cor=TINTA)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17.6), alignment=2)  # 2 = direita
        p.paragraph_format.keep_with_next = True
        _linha_inferior(p)
        for item in grupo["itens"]:
            li = doc.add_paragraph(style="List Bullet")
            li.paragraph_format.space_after = Pt(3)
            li.paragraph_format.line_spacing = 1.2
            _run(li, c.linha_item_memorial(item))


def _finais(doc, memorial) -> None:
    if memorial.financeiro:
        _rotulo_secao(doc, "Condições comerciais")
        _tabela_campos(doc, memorial.financeiro)
    for bloco in memorial.secoes:
        _rotulo_secao(doc, bloco.titulo)
        for trecho in bloco.texto.strip().split("\n"):
            _paragrafo(doc, trecho, depois=3)
    if memorial.contrato:
        _rotulo_secao(doc, "Dados para contrato")
        _tabela_campos(doc, memorial.contrato)
    if memorial.assinatura:
        _paragrafo(doc, "", antes=30, depois=0)
        p = _paragrafo(doc, "_" * 40, depois=2, cor=SUAVE)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _paragrafo(doc, memorial.assinatura[0], depois=1, negrito=True, cor=TINTA)
        for linha in memorial.assinatura[1:]:
            _paragrafo(doc, linha, depois=1, tamanho=9, cor=SUAVE)


def gerar_docx_memorial_descritivo(orcamento, configuracao) -> bytes:
    memorial = c.montar_memorial(orcamento, configuracao)
    doc = Document()
    _pagina(doc)
    _topo(doc, orcamento, configuracao, memorial)
    _tabela_campos(doc, memorial.cabecalho, largura_rotulo=Cm(3.4))
    _paragrafo(doc, memorial.saudacao, antes=14, depois=4)
    _paragrafo(doc, memorial.introducao, depois=6)
    _categorias(doc, memorial.categorias)
    _finais(doc, memorial)

    empresa = c.nome_empresa(configuracao)
    rodape = doc.sections[0].footer.paragraphs[0]
    _run(
        rodape,
        " · ".join(p for p in (empresa, f"Memorial descritivo {orcamento.numero}") if p),
        tamanho=7.5,
        cor=SUAVE,
    )
    doc.core_properties.title = f"Memorial Descritivo {orcamento.numero}"
    doc.core_properties.author = empresa
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
